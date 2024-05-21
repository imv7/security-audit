import os
import docx
import pandas as pd
from openpyxl import Workbook
from openpyxl.utils.dataframe import dataframe_to_rows

def extract_tables_from_docx(docx_file):
    document = docx.Document(docx_file)
    tables_data = []
    for table in document.tables:
        table_data = []
        for row in table.rows:
            row_data = [cell.text.strip() for cell in row.cells]
            table_data.append(row_data)
        tables_data.append(table_data)
    return tables_data

def save_tables_as_spreadsheet(tables_data, output_file):
    wb = Workbook()
    ws = wb.active

    # Add titles for each column
    titles = ["CSD #", "Responsibilities", "Company", "Customer"]
    ws.append(titles)

    # Write tables data to the worksheet
    for table_data in tables_data:
        for row_data in table_data:
            ws.append(row_data)

    # Set column width based on the maximum width of the data in each column
    for col in ws.columns:
        max_length = 0
        column = col[0].column_letter
        for cell in col:
            try:
                if len(str(cell.value)) > max_length:
                    max_length = len(cell.value)
            except:
                pass
        adjusted_width = (max_length + 2) * 1.2
        ws.column_dimensions[column].width = adjusted_width

    # Save the workbook to the output file
    wb.save(output_file)

def filter_data(input_file, output_file):
    # Read the generated Excel file using pandas
    df = pd.read_excel(input_file)

    # Filter the third column (index 2) for "P"
    filtered_df = df[df.iloc[:, 2] == "P"]

    # Save the filtered data back to Excel with titles
    filtered_df.to_excel(output_file, index=False)

def main():
    input_docx_file = "CSD.docx"
    output_excel_file = "CSD.xlsx"
    output_filtered_file = "CSD_Company_Responsibilities.xlsx"

    # Check if the input file exists
    if not os.path.exists(input_docx_file):
        print(f"Error: Input file '{input_docx_file}' not found.")
        return

    # Extract tables from the input docx file
    tables_data = extract_tables_from_docx(input_docx_file)

    # Save the tables data to an Excel file
    save_tables_as_spreadsheet(tables_data, output_excel_file)

    # Filter the data and save it to another Excel file
    filter_data(output_excel_file, output_filtered_file)

    print(f"Kyndryl Responsibilities extracted to {output_filtered_file}")

if __name__ == "__main__":
    main()